"""Organization-scoped policy document management APIs.

Authentication uses the existing user API key (Authorization: Bearer aigw_...).
The caller must have an active organization membership. Only org_admin may upload,
activate, version, or archive documents; org_viewer may list and inspect them.
Uploaded files are stored in a private server-side directory, not served as static files.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, Form, status
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.auth.dependencies import get_current_principal
from app.database import get_db
from app.models import (
    Organization,
    OrganizationMembership,
    PolicyDocument,
    PolicyDocumentVersion,
    Principal,
)

router = APIRouter(prefix="/v1/organizations", tags=["Organization Policy Documents"])

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MiB
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
ALLOWED_ROLES = {"org_admin", "org_viewer"}

PROJECT_ROOT = Path(__file__).resolve().parents[2]
_STORAGE_SETTING = os.getenv("POLICY_STORAGE_DIR")
STORAGE_ROOT = (
    Path(_STORAGE_SETTING).expanduser().resolve()
    if _STORAGE_SETTING
    else (PROJECT_ROOT / "private_storage" / "policy_documents").resolve()
)


def _require_membership(
    db: Session,
    principal: Principal,
    organization_id: int,
    *,
    admin_only: bool = False,
) -> OrganizationMembership:
    """Authorize using the authenticated principal and persisted org membership."""
    if not principal.is_active:
        raise HTTPException(status_code=403, detail="This principal is inactive.")

    organization = db.get(Organization, organization_id)
    if organization is None or not organization.is_active:
        raise HTTPException(status_code=404, detail="Organization not found or inactive.")

    membership = (
        db.query(OrganizationMembership)
        .filter(
            OrganizationMembership.organization_id == organization_id,
            OrganizationMembership.principal_id == principal.id,
            OrganizationMembership.is_active.is_(True),
        )
        .first()
    )
    if membership is None:
        raise HTTPException(
            status_code=403,
            detail="You do not have an active membership in this organization.",
        )

    if membership.role not in ALLOWED_ROLES:
        raise HTTPException(status_code=403, detail="Your organization role is not permitted.")

    # Defence in depth: the membership must agree with the principal's current team.
    team_org_id = principal.team.organization_id if principal.team else None
    if team_org_id != organization_id:
        raise HTTPException(
            status_code=403,
            detail="Your principal's team is not assigned to this organization.",
        )

    if admin_only and membership.role != "org_admin":
        raise HTTPException(
            status_code=403,
            detail="The org_admin role is required for this operation.",
        )

    return membership


def _safe_original_filename(filename: str | None) -> str:
    # Strip both POSIX and Windows path components, then remove control/path characters.
    base = (filename or "upload").replace("\\", "/").split("/")[-1].strip()
    base = re.sub(r"[^\w.() -]", "_", base, flags=re.UNICODE)
    return (base or "upload")[:255]


async def _read_upload_limited(upload: UploadFile) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await upload.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_FILE_SIZE_BYTES:
            raise HTTPException(
                status_code=413,
                detail="File exceeds the 10 MiB upload limit.",
            )
        chunks.append(chunk)

    content = b"".join(chunks)
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")
    return content


def _extract_text(extension: str, content: bytes) -> str:
    """Extract searchable text and add page/heading markers for later citations."""
    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise HTTPException(status_code=415, detail="The file does not look like a valid PDF.")
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            pages: list[str] = []
            for page_number, page in enumerate(reader.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    pages.append(f"[PAGE {page_number}]\n{page_text}")
            extracted = "\n\n".join(pages)
        elif extension == ".docx":
            if not content.startswith(b"PK\x03\x04"):
                raise HTTPException(status_code=415, detail="The file does not look like a valid DOCX document.")
            from docx import Document as WordDocument

            doc = WordDocument(io.BytesIO(content))
            paragraphs: list[str] = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = getattr(paragraph.style, "name", "") or ""
                if style_name.lower().startswith("heading"):
                    paragraphs.append(f"\n[SECTION: {text}]\n")
                else:
                    paragraphs.append(text)
            extracted = "\n".join(paragraphs).strip()
        else:
            raise HTTPException(status_code=415, detail="Only PDF and DOCX files are supported.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="The document could not be parsed. Confirm it is a valid, unencrypted PDF or DOCX.",
        ) from exc

    if not extracted.strip():
        if extension == ".pdf":
            detail = (
                "No selectable text was found in this PDF. Scanned PDFs require OCR, "
                "which is not enabled in this implementation."
            )
        else:
            detail = "No readable text was found in this DOCX file."
        raise HTTPException(status_code=422, detail=detail)
    return extracted


def _document_summary(document: PolicyDocument) -> dict:
    versions = sorted(document.versions or [], key=lambda v: v.version_number)
    return {
        "id": document.id,
        "scope": document.scope,
        "organization_id": document.organization_id,
        "title": document.title,
        "status": document.status,
        "created_by_principal_id": document.created_by_principal_id,
        "created_at": document.created_at,
        "updated_at": document.updated_at,
        "versions": [
            {
                "id": version.id,
                "version_number": version.version_number,
                "original_filename": version.original_filename,
                "content_type": version.content_type,
                "file_size_bytes": version.file_size_bytes,
                "sha256": version.sha256,
                "processing_status": version.processing_status,
                "is_active": version.is_active,
                "uploaded_by_principal_id": version.uploaded_by_principal_id,
                "created_at": version.created_at,
            }
            for version in versions
        ],
    }


def _get_org_document(db: Session, organization_id: int, document_id: int) -> PolicyDocument:
    document = (
        db.query(PolicyDocument)
        .filter(
            PolicyDocument.id == document_id,
            PolicyDocument.scope == "organization",
            PolicyDocument.organization_id == organization_id,
        )
        .first()
    )
    if document is None:
        raise HTTPException(status_code=404, detail="Policy document not found in this organization.")
    return document


def _save_private_file(organization_id: int, extension: str, content: bytes) -> Path:
    """Persist to an opaque, private filename; never use a client path on disk."""
    folder = STORAGE_ROOT / f"org_{organization_id}"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{uuid4().hex}{extension}"
    path.write_bytes(content)
    return path.resolve()


@router.post("/{organization_id}/policy-documents", status_code=status.HTTP_201_CREATED)
async def upload_policy_document(
    organization_id: int,
    title: str = Form(..., min_length=2, max_length=255),
    file: UploadFile = File(...),
    principal: Principal = Depends(get_current_principal),
    db: Session = Depends(get_db),
):
    """Create an organization policy and its first, inactive version."""
    _require_membership(db, principal, organization_id, admin_only=True)

    original_filename = _safe_original_filename(file.filename)
    extension = Path(original_filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Only PDF and DOCX files are supported.")

    content = await _read_upload_limited(file)
    extracted_text = _extract_text(extension, content)
    canonical_content_type = (
        "application/pdf"
        if extension == ".pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    private_path: Path | None = None
    try:
        document = PolicyDocument(
            scope="organization",
            organization_id=organization_id,
            title=title.strip(),
            status="draft",
            created_by_principal_id=principal.id,
        )
        db.add(document)
        db.flush()  # Allocate document ID before committing the related version.

        private_path = _save_private_file(organization_id, extension, content)
        version = PolicyDocumentVersion(
            document_id=document.id,
            version_number=1,
            original_filename=original_filename,
            content_type=canonical_content_type,
            file_size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
            storage_path=str(private_path),
            extracted_text=extracted_text,
            processing_status="ready",
            is_active=False,
            uploaded_by_principal_id=principal.id,
        )
        db.add(version)
        db.commit()
        db.refresh(document)
        return _document_summary(document)
    except Exception:
        db.rollback()
        if private_path is not None:
            try:
                private_path.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    finally:
        await file.close()


@router.get("/{organization_id}/policy-documents")
def list_policy_documents(
    organization_id: int,
    principal: Principal = Depends(get_current_principal),
    db: Session = Depends(get_db),
):
    """List documents belonging only to the authenticated caller's organization."""
    _require_membership(db, principal, organization_id)
    documents = (
        db.query(PolicyDocument)
        .filter(
            PolicyDocument.scope == "organization",
            PolicyDocument.organization_id == organization_id,
        )
        .order_by(PolicyDocument.updated_at.desc(), PolicyDocument.id.desc())
        .all()
    )
    return [_document_summary(document) for document in documents]


@router.get("/{organization_id}/policy-documents/{document_id}")
def get_policy_document(
    organization_id: int,
    document_id: int,
    principal: Principal = Depends(get_current_principal),
    db: Session = Depends(get_db),
):
    _require_membership(db, principal, organization_id)
    document = _get_org_document(db, organization_id, document_id)
    return _document_summary(document)


@router.post(
    "/{organization_id}/policy-documents/{document_id}/versions",
    status_code=status.HTTP_201_CREATED,
)
async def upload_policy_document_version(
    organization_id: int,
    document_id: int,
    file: UploadFile = File(...),
    principal: Principal = Depends(get_current_principal),
    db: Session = Depends(get_db),
):
    """Upload a new inactive version; the currently active version remains live."""
    _require_membership(db, principal, organization_id, admin_only=True)
    document = _get_org_document(db, organization_id, document_id)
    if document.status == "archived":
        raise HTTPException(status_code=409, detail="Archived documents cannot receive new versions.")

    original_filename = _safe_original_filename(file.filename)
    extension = Path(original_filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Only PDF and DOCX files are supported.")

    content = await _read_upload_limited(file)
    extracted_text = _extract_text(extension, content)
    canonical_content_type = (
        "application/pdf"
        if extension == ".pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    private_path: Path | None = None
    try:
        latest_version = (
            db.query(func.max(PolicyDocumentVersion.version_number))
            .filter(PolicyDocumentVersion.document_id == document.id)
            .scalar()
        ) or 0
        private_path = _save_private_file(organization_id, extension, content)
        version = PolicyDocumentVersion(
            document_id=document.id,
            version_number=int(latest_version) + 1,
            original_filename=original_filename,
            content_type=canonical_content_type,
            file_size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
            storage_path=str(private_path),
            extracted_text=extracted_text,
            processing_status="ready",
            is_active=False,
            uploaded_by_principal_id=principal.id,
        )
        # Keep the document active if an older active version is still in use.
        # The newly uploaded version remains inactive until explicitly activated.
        db.add(version)
        db.commit()
        db.refresh(document)
        return _document_summary(document)
    except Exception:
        db.rollback()
        if private_path is not None:
            try:
                private_path.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    finally:
        await file.close()


@router.post("/{organization_id}/policy-documents/{document_id}/versions/{version_id}/activate")
def activate_policy_document_version(
    organization_id: int,
    document_id: int,
    version_id: int,
    principal: Principal = Depends(get_current_principal),
    db: Session = Depends(get_db),
):
    """Activate a ready version; only one version per document can be active."""
    _require_membership(db, principal, organization_id, admin_only=True)
    document = _get_org_document(db, organization_id, document_id)
    version = (
        db.query(PolicyDocumentVersion)
        .filter(
            PolicyDocumentVersion.id == version_id,
            PolicyDocumentVersion.document_id == document.id,
        )
        .first()
    )
    if version is None:
        raise HTTPException(status_code=404, detail="Policy document version not found.")
    if version.processing_status != "ready" or not version.extracted_text:
        raise HTTPException(status_code=409, detail="Only successfully processed versions can be activated.")

    for previous in document.versions:
        previous.is_active = previous.id == version.id
    document.status = "active"
    db.commit()
    db.refresh(document)
    return _document_summary(document)


@router.post("/{organization_id}/policy-documents/{document_id}/archive")
def archive_policy_document(
    organization_id: int,
    document_id: int,
    principal: Principal = Depends(get_current_principal),
    db: Session = Depends(get_db),
):
    """Archive an organization policy without deleting audit/history data."""
    _require_membership(db, principal, organization_id, admin_only=True)
    document = _get_org_document(db, organization_id, document_id)
    document.status = "archived"
    for version in document.versions:
        version.is_active = False
    db.commit()
    db.refresh(document)
    return _document_summary(document)
